import os, re, io, csv, math, html
from flask import Flask, request, jsonify, send_file, send_from_directory
from youtube_transcript_api import YouTubeTranscriptApi
from requests import Session
from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

app = Flask(__name__, static_folder='public', static_url_path='')
app.config['MAX_CONTENT_LENGTH'] = 8 * 1024 * 1024
ID = re.compile(r'^[A-Za-z0-9_-]{11}$')

class TimedSession(Session):
    def request(self, *args, **kwargs):
        kwargs.setdefault('timeout', (5, 20))
        return super().request(*args, **kwargs)

def stamp(seconds, sep='.'):
    ms = round(max(0, seconds) * 1000)
    h, ms = divmod(ms, 3600000); m, ms = divmod(ms, 60000); s, ms = divmod(ms, 1000)
    return f'{h:02}:{m:02}:{s:02}{sep}{ms:03}'

def validate_cues(raw):
    if not isinstance(raw, list) or not raw or len(raw) > 50000:
        raise ValueError('Provide 1–50,000 timestamped transcript lines.')
    cues = []
    for c in raw:
        start, duration = float(c['start']), float(c['duration'])
        if not math.isfinite(start) or not math.isfinite(duration) or start < 0 or duration < 0 or start > 100000000 or duration > 100000000:
            raise ValueError('Invalid transcript timing.')
        text = str(c['text'])[:20000]
        cues.append(dict(start=start, duration=duration, text=text))
    return sorted(cues, key=lambda c: c['start'])

@app.get('/')
def index(): return send_from_directory('public', 'index.html')

@app.get('/health')
def health(): return {'status': 'ok', 'application': 'YouTube Playlist Transcript Player'}

@app.get('/api/transcript/<video_id>')
def transcript(video_id):
    if not ID.fullmatch(video_id):
        return jsonify(error='Invalid YouTube video ID.'), 400

    language = request.args.get('language', 'en')
    if not re.fullmatch(r'[a-zA-Z-]{2,15}', language):
        return jsonify(error='Invalid language code.'), 400

    key = os.environ.get('SUPADATA_API_KEY', '').strip()
    if not key:
        return jsonify(error='Add SUPADATA_API_KEY in Render Environment.'), 503

    try:
        with TimedSession() as session:
            response = session.get(
                'https://api.supadata.ai/v1/transcript',
                params={
                    'url': f'https://www.youtube.com/watch?v={video_id}',
                    'mode': 'native',
                    'text': 'false',
                    'lang': language
                },
                headers={'x-api-key': key},
                allow_redirects=False
            )

        messages = {
            206: 'No existing captions are available. AI generation is disabled.',
            202: 'Unexpected background job. No AI generation was requested.',
            401: 'Supadata rejected the API key. Check Render Environment.',
            402: 'Supadata says this request requires a different plan.',
            403: 'Supadata cannot access this video. It may be restricted.',
            404: 'Supadata could not find an accessible video.',
            429: 'Supadata rate or credit limit reached. Check your usage.'
        }

        if response.status_code in messages:
            status = 429 if response.status_code == 429 else 422
            return jsonify(error=messages[response.status_code]), status

        if response.status_code != 200:
            return jsonify(error='The transcript provider returned an error.'), 502

        data = response.json()
        rows = data.get('content')
        if not isinstance(rows, list) or not rows:
            return jsonify(error='No timestamped captions were returned.'), 422

        cues = []
        for row in rows:
            if not isinstance(row.get('text'), str):
                raise ValueError('Invalid caption text')

            for field in ('offset', 'duration'):
                value = row.get(field)
                if (
                    isinstance(value, bool)
                    or not isinstance(value, (int, float))
                    or not math.isfinite(value)
                    or value < 0
                ):
                    raise ValueError('Invalid caption timing')

            cues.append({
                'start': row['offset'] / 1000,
                'duration': row['duration'] / 1000,
                'text': re.sub(r'\s+', ' ', row['text']).strip()
            })

        return jsonify(
            videoId=video_id,
            language=data.get('lang', language),
            provider='Supadata',
            cues=cues
        )

    except Exception:
        return jsonify(error='Transcript retrieval failed. Please try again.'), 502
@app.post('/api/export')
def export():
    try:
        body = request.get_json(); cues = validate_cues(body['cues'])
        fmt = body['format'].lower(); title = str(body.get('title', 'Transcript'))[:200]
        minutes = int(body.get('minutes', 0))
        if minutes not in (0, 5, 10, 15, 30, 60): raise ValueError('Invalid segment size.')
        if fmt not in ('txt','md','docx','pdf','csv','srt','vtt'): raise ValueError('Unsupported export format.')
    except (ValueError, TypeError, KeyError, AttributeError):
        return jsonify(error='Invalid export request.'), 400
    out = io.BytesIO(); mime = 'text/plain; charset=utf-8'
    groups = {}
    for c in cues:
        n = int(c['start'] // (minutes * 60)) if minutes else 0
        groups.setdefault(n, []).append(c)
    def heading(n): return f'Segment {n+1}: {stamp(n*minutes*60)[:-4]} – {stamp((n+1)*minutes*60)[:-4]}'
    lines = [title, '']
    for n, group in groups.items():
        if minutes: lines.extend([heading(n), ''])
        lines.extend(f"[{stamp(c['start'])}] {c['text']}" for c in group)
        lines.append('')
    if fmt in ('txt', 'md'):
        if fmt == 'md': lines[0] = '# ' + title
        out.write('\n'.join(lines).encode('utf-8'))
    elif fmt in ('srt','vtt'):
        text = 'WEBVTT\n\n' if fmt == 'vtt' else ''
        sep = ',' if fmt == 'srt' else '.'
        for i,c in enumerate(cues,1):
            text += f"{i}\n{stamp(c['start'],sep)} --> {stamp(c['start']+c['duration'],sep)}\n{c['text']}\n\n"
        out.write(text.encode('utf-8')); mime = 'text/vtt' if fmt == 'vtt' else 'application/x-subrip'
    elif fmt == 'csv':
        stream=io.StringIO(); writer=csv.writer(stream); writer.writerow(['segment','start_seconds','duration_seconds','timestamp','text'])
        for c in cues:
            value=c['text']
            if value.lstrip().startswith(('=', '+', '-', '@')): value="'"+value
            writer.writerow([int(c['start']//(minutes*60))+1 if minutes else 1,c['start'],c['duration'],stamp(c['start']),value])
        out.write(stream.getvalue().encode('utf-8-sig')); mime='text/csv'
    elif fmt == 'docx':
        doc=Document(); doc.add_heading(title,0)
        for n,group in groups.items():
            if minutes: doc.add_heading(heading(n),1)
            for c in group: doc.add_paragraph(f"[{stamp(c['start'])}] {c['text']}")
        doc.save(out); mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif fmt == 'pdf':
        styles=getSampleStyleSheet()
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        font=os.path.join(os.path.dirname(__file__),'fonts','DejaVuSans.ttf')
        if os.path.exists(font):
            if 'Transcript' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont('Transcript',font))
            for style in styles.byName.values(): style.fontName='Transcript'
        story=[Paragraph(html.escape(title), styles['Title']),Spacer(1,12)]
        for n,group in groups.items():
            if minutes: story.append(Paragraph(heading(n),styles['Heading2']))
            for c in group:
                story.extend([Paragraph(html.escape(f"[{stamp(c['start'])}] {c['text']}"),styles['BodyText']),Spacer(1,6)])
        SimpleDocTemplate(out).build(story); mime='application/pdf'
    out.seek(0)
    safe=re.sub(r'[^A-Za-z0-9_-]+','-',title).strip('-')[:80] or 'transcript'
    return send_file(out,mimetype=mime,as_attachment=True,download_name=f'{safe}.{fmt}')

@app.errorhandler(413)
def too_large(e): return jsonify(error='Request exceeds the 8 MB limit.'),413

if __name__ == '__main__': app.run(host='0.0.0.0',port=int(os.environ.get('PORT',3000)))
